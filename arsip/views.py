import zipfile
import io
from pathlib import Path
from datetime import datetime

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from django.db.models import Q

from .models import Arsip, Folder
from .forms import ArsipForm, BulkArsipForm, FolderUploadForm, StaffCreationForm
from .utils import build_breadcrumb, build_full_tree
from django.shortcuts import redirect
from .models import Arsip, Folder, ActivityLog


def log_activity(user, action, description, target_type='', target_id=None, request=None):
    """Helper buat nyatet aktivitas"""
    ip = None
    if request:
        x_forwarded = request.META.get('HTTP_X_FORWARDED_FOR')
        if x_forwarded:
            ip = x_forwarded.split(',')[0]
        else:
            ip = request.META.get('REMOTE_ADDR')
    
    try:
        ActivityLog.objects.create(
            user=user if user.is_authenticated else None,
            action=action,
            description=description,
            target_type=target_type,
            target_id=target_id,
            ip_address=ip,
        )
    except Exception as e:
        print(f"[log_activity] Error: {e}")

def is_admin(user):
    """Cek user punya akses admin"""
    return user.is_authenticated and (user.is_staff or user.is_superuser)

def dashboard_redirect(request):
    """Root URL: admin → dashboard, lainnya → public home"""
    if request.user.is_authenticated and request.user.is_staff:
        return redirect('arsip:list')
    return redirect('arsip:public_home')


# ============================================
# ARSIP LIST & DETAIL
# ============================================

def arsip_list(request):
    # User biasa (bukan admin) → redirect ke public home
    if request.user.is_authenticated and not request.user.is_staff:
        return redirect('arsip:public_home')
    """Dashboard admin + daftar arsip dengan pagination"""
    from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
    from django.db.models import Sum, Count
    
    query = request.GET.get('q', '')
    tahun = request.GET.get('tahun', '')
    kategori = request.GET.get('kategori', '')
    page = request.GET.get('page', 1)

    base_qs = Arsip.objects.select_related('folder').filter(
        is_active=True, folder__is_active=True
    )

    arsip_qs = base_qs
    if query:
        arsip_qs = arsip_qs.filter(
            Q(name__icontains=query) |
            Q(description__icontains=query) |
            Q(creator__icontains=query) |
            Q(reference_code__icontains=query)
        )
    if tahun:
        arsip_qs = arsip_qs.filter(folder__year=tahun)
    if kategori:
        arsip_qs = arsip_qs.filter(file_type=kategori)

    # Pagination — 20 per halaman
    paginator = Paginator(arsip_qs, 20)
    try:
        arsip_page = paginator.page(page)
    except PageNotAnInteger:
        arsip_page = paginator.page(1)
    except EmptyPage:
        arsip_page = paginator.page(paginator.num_pages)

    tahun_list = Folder.objects.filter(
        is_active=True
    ).values_list('year', flat=True).distinct().order_by('-year')

    # Statistik
    total_arsip = base_qs.count()
    total_folder = Folder.objects.filter(is_active=True).count()
    total_size = base_qs.aggregate(total=Sum('size_bytes'))['total'] or 0
    
    pdf_count = base_qs.filter(file_type='pdf').count()
    xlsx_count = base_qs.filter(file_type='xlsx').count()
    docx_count = base_qs.filter(file_type='docx').count()
    lain_count = base_qs.filter(file_type='lain').count()
    
    level_fonds = Folder.objects.filter(level='fonds', is_active=True).count()
    level_subfonds = Folder.objects.filter(level='subfonds', is_active=True).count()
    level_series = Folder.objects.filter(level='series', is_active=True).count()
    level_file = Folder.objects.filter(level='file', is_active=True).count()
    
    arsip_per_tahun = base_qs.values('folder__year').annotate(
        count=Count('id')
    ).order_by('-folder__year')[:10]
    
    recent_arsip = base_qs.order_by('-uploaded_at')[:5]

    context = {
        'arsip_list': arsip_page,
        'alls_folders': Folder.objects.filter(is_active=True).order_by('level', 'name'),
        'page_obj': arsip_page,
        'paginator': paginator,
        'total_hasil': paginator.count,
        'tahun_list': tahun_list,
        'query': query,
        'tahun_aktif': tahun,
        'kategori_aktif': kategori,
        'total_arsip': total_arsip,
        'total_folder': total_folder,
        'total_size': total_size,
        'pdf_count': pdf_count,
        'xlsx_count': xlsx_count,
        'docx_count': docx_count,
        'lain_count': lain_count,
        'level_fonds': level_fonds,
        'level_subfonds': level_subfonds,
        'level_series': level_series,
        'level_file': level_file,
        'arsip_per_tahun': arsip_per_tahun,
        'recent_arsip': recent_arsip,
    }
    if request.user.is_authenticated and request.user.is_staff:
        template = 'arsip/admin_dashboard.html'
    else:
        template = 'arsip/list.html'

    return render(request, template, context)

    # ============================================
    # STATISTIK DASHBOARD
    # ============================================
    total_arsip = base_qs.count()
    total_folder = Folder.objects.filter(is_active=True).count()
    
    # Total size (bytes)
    from django.db.models import Sum
    total_size = base_qs.aggregate(total=Sum('size_bytes'))['total'] or 0
    
    # Breakdown per tipe file
    pdf_count = base_qs.filter(file_type='pdf').count()
    xlsx_count = base_qs.filter(file_type='xlsx').count()
    docx_count = base_qs.filter(file_type='docx').count()
    lain_count = base_qs.filter(file_type='lain').count()
    
    # Breakdown per level folder
    level_fonds = Folder.objects.filter(level='fonds', is_active=True).count()
    level_subfonds = Folder.objects.filter(level='subfonds', is_active=True).count()
    level_series = Folder.objects.filter(level='series', is_active=True).count()
    level_file = Folder.objects.filter(level='file', is_active=True).count()
    
    # Arsip per tahun (untuk grafik)
    from django.db.models import Count
    arsip_per_tahun = base_qs.values('folder__year').annotate(
        count=Count('id')
    ).order_by('-folder__year')[:10]
    
    # Recent arsip (5 terbaru)
    recent_arsip = base_qs.order_by('-uploaded_at')[:5]

    context = {
        'arsip_list': arsip_qs,
        'tahun_list': tahun_list,
        'query': query,
        'tahun_aktif': tahun,
        'kategori_aktif': kategori,
        
        # Statistik
        'total_arsip': total_arsip,
        'total_folder': total_folder,
        'total_size': total_size,
        'pdf_count': pdf_count,
        'xlsx_count': xlsx_count,
        'docx_count': docx_count,
        'lain_count': lain_count,
        'level_fonds': level_fonds,
        'level_subfonds': level_subfonds,
        'level_series': level_series,
        'level_file': level_file,
        'arsip_per_tahun': arsip_per_tahun,
        'recent_arsip': recent_arsip,
    }
    return render(request, 'arsip/list.html', context)


def arsip_detail(request, pk):
    """Halaman detail 1 arsip"""
    arsip = get_object_or_404(Arsip, pk=pk, is_active=True)
    return render(request, 'arsip/detail.html', {'arsip': arsip})


# ============================================
# ARSIP EDIT & DELETE
# ============================================

@login_required
def arsip_edit(request, pk):
    """Edit arsip: nama + metadata 6 area + pindah folder"""
    arsip = get_object_or_404(Arsip, pk=pk)
    
    if request.method == 'POST':
        arsip.name = request.POST.get('name', arsip.name)
        arsip.file_type = request.POST.get('file_type', arsip.file_type)
        arsip.reference_code = request.POST.get('reference_code', '')
        arsip.creator = request.POST.get('creator', '')
        arsip.repository = request.POST.get('repository', '')
        arsip.description = request.POST.get('description', '')
        arsip.scope_content = request.POST.get('scope_content', '')
        arsip.note = request.POST.get('note', '')
        arsip.subject_access = request.POST.get('subject_access', '')
        arsip.location_physical = request.POST.get('location_physical', '')
        arsip.access_level = request.POST.get('access_level', 'Terbatas')
        arsip.media_type = request.POST.get('media_type', '')
        arsip.extent_medium = request.POST.get('extent_medium', '')
        arsip.level_description = request.POST.get('level_description', 'item')
        arsip.tags = request.POST.get('tags', '')
        
        date_str = request.POST.get('date_created', '')
        if date_str:
            try:
                arsip.date_created = datetime.strptime(date_str, '%Y-%m-%d').date()
            except ValueError:
                pass
        
        folder_id = request.POST.get('folder', '')
        if folder_id:
            try:
                arsip.folder = Folder.objects.get(pk=folder_id)
            except Folder.DoesNotExist:
                pass
        
        arsip.save()

        log_activity(
            request.user, 'edit',
            f'Edit arsip "{arsip.name}"',
            target_type='arsip', target_id=arsip.pk,
            request=request
        )

        messages.success(request, f'Arsip "{arsip.name}" berhasil diupdate.')
        return redirect('arsip:detail', pk=arsip.pk)
    
    all_folders = Folder.objects.filter(is_active=True).order_by('level', 'name')
    return render(request, 'arsip/arsip_edit.html', {
        'arsip': arsip,
        'all_folders': all_folders,
    })


@login_required
def arsip_delete(request, pk):
    """Hapus arsip"""
    arsip = get_object_or_404(Arsip, pk=pk)

    if request.method == 'POST':
        # Simpen dulu datanya SEBELUM delete
        nama = arsip.name
        arsip_id = arsip.pk
        
        # Baru delete
        arsip.delete()
        
        # Log aktivitas
        log_activity(
            request.user, 'delete',
            f'Hapus arsip "{nama}"',
            target_type='arsip', target_id=arsip_id,
            request=request
        )
        
        messages.success(request, f'Arsip "{nama}" berhasil dihapus.')
        return redirect('arsip:list')

    return render(request, 'arsip/delete.html', {'arsip': arsip})


# ============================================
# UPLOAD
# ============================================

@user_passes_test(is_admin)
def arsip_upload(request):
    """Upload arsip tunggal"""
    if request.method == 'POST':
        form = ArsipForm(request.POST, request.FILES)
        if form.is_valid():
            arsip = form.save(commit=False)
            arsip.size_bytes = request.FILES['file'].size
            arsip.save()

            log_activity(
                request.user, 'upload',
                f'Upload file "{arsip.name}"',
                target_type='arsip', target_id=arsip.pk,
                request=request
            )

            messages.success(request, f'Arsip "{arsip.name}" berhasil diunggah.')
            return redirect('arsip:list')
    else:
        form = ArsipForm()
    return render(request, 'arsip/upload.html', {'form': form})


@login_required
def bulk_upload(request):
    """Bulk upload banyak file ke 1 folder"""
    if request.method == 'POST':
        form = BulkArsipForm(request.POST, request.FILES)
        if form.is_valid():
            folder = form.cleaned_data['folder']
            files = form.cleaned_data['files']
            file_type = form.cleaned_data['file_type']
            creator = form.cleaned_data['creator']
            access_level = form.cleaned_data['access_level']
            
            success = 0
            errors = []
            for f in files:
                try:
                    arsip = Arsip(
                        name=f.name.rsplit('.', 1)[0],
                        folder=folder,
                        file_type=file_type,
                        size_bytes=f.size,
                        creator=creator,
                        access_level=access_level,
                    )
                    arsip.file.save(f.name, f, save=True)
                    success += 1
                except Exception as e:
                    errors.append(f"{f.name}: {e}")
            
            if success:
                messages.success(request, f'{success} file berhasil diupload ke "{folder.name}".')
            if errors:
                messages.error(request, f'{len(errors)} file gagal.')
            return redirect('arsip:folder', pk=folder.pk)
    else:
        initial = {}
        folder_id = request.GET.get('folder')
        if folder_id:
            initial['folder'] = folder_id
        form = BulkArsipForm(initial=initial)
    return render(request, 'arsip/bulk_upload.html', {'form': form})


@login_required
def folder_upload(request):
    """Upload ZIP, auto-extract struktur folder"""
    if request.method == 'POST':
        form = FolderUploadForm(request.POST, request.FILES)
        if form.is_valid():
            zip_file = request.FILES['zip_file']
            root_name = form.cleaned_data['root_fonds']
            year = form.cleaned_data['year']
            
            fonds, _ = Folder.objects.get_or_create(
                name=root_name, parent=None,
                defaults={'level': 'fonds', 'year': year}
            )
            
            success_files = 0
            success_folders = 0
            errors = []
            
            try:
                with zipfile.ZipFile(zip_file, 'r') as z:
                    for member in z.namelist():
                        if member.endswith('/'):
                            continue
                        if member.startswith('__MACOSX') or '/.' in member:
                            continue
                        
                        parts = Path(member).parts
                        if len(parts) < 2:
                            continue
                        
                        parent = fonds
                        level_map = ['subfonds', 'series', 'file']
                        for i, part in enumerate(parts[:-1]):
                            level = level_map[min(i, len(level_map)-1)]
                            parent, created = Folder.objects.get_or_create(
                                name=part, parent=parent,
                                defaults={'level': level, 'year': year}
                            )
                            if created:
                                success_folders += 1
                        
                        try:
                            file_data = z.read(member)
                            file_name = parts[-1]
                            ext = file_name.rsplit('.', 1)[-1].lower() if '.' in file_name else ''
                            
                            if ext in ('xlsx', 'xls'):
                                file_type = 'xlsx'
                            elif ext in ('docx', 'doc'):
                                file_type = 'docx'
                            elif ext == 'pdf':
                                file_type = 'pdf'
                            else:
                                file_type = 'lain'
                            
                            arsip = Arsip(
                                name=file_name.rsplit('.', 1)[0] if '.' in file_name else file_name,
                                folder=parent,
                                file_type=file_type,
                                size_bytes=len(file_data),
                            )
                            arsip.file.save(file_name, io.BytesIO(file_data), save=True)
                            success_files += 1
                        except Exception as e:
                            errors.append(f"{member}: {e}")
            
            except zipfile.BadZipFile:
                messages.error(request, 'File bukan ZIP yang valid.')
                return render(request, 'arsip/folder_upload.html', {'form': form})
            
            messages.success(request, f'Berhasil! {success_folders} folder + {success_files} file diupload.')
            if errors:
                messages.error(request, f'{len(errors)} error.')
            return redirect('arsip:folder', pk=fonds.pk)
    else:
        form = FolderUploadForm()
    return render(request, 'arsip/folder_upload.html', {'form': form})


# ============================================
# FOLDER BROWSE & EDIT
# ============================================

def folder_browse(request, pk=None):
    """Browse folder (tree view)"""
    if pk:
        current_folder = get_object_or_404(Folder, pk=pk, is_active=True)
    else:
        current_folder = None

    if current_folder:
        subfolders = Folder.objects.filter(parent=current_folder, is_active=True).order_by('order', 'name')
        arsip_list = Arsip.objects.filter(folder=current_folder, is_active=True).order_by('-uploaded_at')
    else:
        subfolders = Folder.objects.filter(parent__isnull=True, is_active=True).order_by('order', 'name')
        arsip_list = Arsip.objects.none()

    breadcrumb = build_breadcrumb(current_folder) if current_folder else []
    full_tree = build_full_tree()

    # ============================================
    # HITUNG ANCESTOR IDs (parent + grandparent + dst)
    # ============================================
    ancestor_ids = []
    if current_folder:
        # Folder aktif sendiri
        ancestor_ids.append(current_folder.pk)
        # Semua parent ke atas
        parent = current_folder.parent
        while parent:
            ancestor_ids.append(parent.pk)
            parent = parent.parent

    context = {
        'current_folder': current_folder,
        'subfolders': subfolders,
        'arsip_list': arsip_list,
        'breadcrumb': breadcrumb,
        'full_tree': full_tree,
        'current_folder_id': current_folder.id if current_folder else None,
        'ancestor_ids': ancestor_ids,  # ← BARU
        'all_folders': Folder.objects.filter(is_active=True).order_by('level', 'name'),
    }
    return render(request, 'arsip/folder_browse.html', context)


@login_required
def folder_edit(request, pk):
    """Edit folder: rename, move, change level"""
    folder = get_object_or_404(Folder, pk=pk)
    
    if request.method == 'POST':
        folder.name = request.POST.get('name', folder.name)
        folder.level = request.POST.get('level', folder.level)
        folder.year = request.POST.get('year', folder.year)
        folder.reference_code = request.POST.get('reference_code', '')
        folder.description = request.POST.get('description', '')
        folder.tags = request.POST.get('tags', '')
        
        try:
            folder.order = int(request.POST.get('order', 0))
        except ValueError:
            folder.order = 0
        
        parent_id = request.POST.get('parent', '')
        if parent_id:
            try:
                folder.parent = Folder.objects.get(pk=parent_id)
            except Folder.DoesNotExist:
                folder.parent = None
        else:
            folder.parent = None
        
        folder.save()

        log_activity(
            request.user, 'edit',
            f'Edit folder "{folder.name}"',
            target_type='folder', target_id=folder.pk,
            request=request
        )

        messages.success(request, f'Folder "{folder.name}" berhasil diupdate.')
        return redirect('arsip:folder', pk=folder.pk)
    
    all_folders = Folder.objects.filter(is_active=True).exclude(pk=folder.pk).order_by('level', 'name')
    return render(request, 'arsip/folder_edit.html', {
        'folder': folder,
        'all_folders': all_folders,
    })


# ============================================
# USER MANAGEMENT
# ============================================

@user_passes_test(is_admin)
def user_list(request):
    users = User.objects.all().order_by('username')
    return render(request, 'arsip/user_list.html', {'users': users})


@user_passes_test(is_admin)
def user_create(request):
    if request.method == 'POST':
        form = StaffCreationForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            role = form.cleaned_data.get('role', 'user')
            user.is_staff = (role == 'admin')
            user.is_superuser = False
            user.save()
            messages.success(request, f'Akun "{user.username}" ({role}) berhasil dibuat.')
            return redirect('arsip:user_list')
    else:
        form = StaffCreationForm()
    return render(request, 'arsip/user_form.html', {'form': form})


@user_passes_test(is_admin)
def user_toggle_active(request, pk):
    user = get_object_or_404(User, pk=pk)
    if user == request.user:
        messages.error(request, 'Gak bisa nonaktifin akun sendiri.')
        return redirect('arsip:user_list')
    user.is_active = not user.is_active
    user.save()
    status = 'diaktifkan' if user.is_active else 'dinonaktifkan'
    messages.success(request, f'Akun "{user.username}" {status}.')
    return redirect('arsip:user_list')


@user_passes_test(is_admin)
def user_reset_password(request, pk):
    import secrets
    import string
    user = get_object_or_404(User, pk=pk)
    
    if request.method == 'POST':
        alphabet = string.ascii_letters + string.digits
        new_password = ''.join(secrets.choice(alphabet) for _ in range(10))
        user.set_password(new_password)
        user.save()
        messages.success(request, f'Password baru "{user.username}": {new_password}  (CATAT!)')
        return redirect('arsip:user_list')
    
    return render(request, 'arsip/user_reset.html', {'target_user': user})

def public_home(request):
    """Dashboard publik — bisa diakses tanpa login"""
    total_arsip = Arsip.objects.filter(is_active=True, folder__is_active=True).count()
    total_folder = Folder.objects.filter(is_active=True).count()
    recent_arsip = Arsip.objects.filter(
        is_active=True, folder__is_active=True
    ).order_by('-uploaded_at')[:8]
    
    full_tree = build_full_tree()
    
    context = {
        'total_arsip': total_arsip,
        'total_folder': total_folder,
        'recent_arsip': recent_arsip,
        'full_tree': full_tree,
    }
    return render(request, 'arsip/public_home.html', context)

def custom_404(request, exception):
    """Handler 404 error"""
    return render(request, '404.html', status=404)


def custom_500(request):
    """Handler 500 error"""
    return render(request, '500.html', status=500)


@user_passes_test(is_admin)
def export_excel(request):
    """Export daftar arsip ke Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from django.http import HttpResponse
    from datetime import datetime

    query = request.GET.get('q', '')
    tahun = request.GET.get('tahun', '')
    kategori = request.GET.get('kategori', '')

    arsip_qs = Arsip.objects.select_related('folder').filter(
        is_active=True, folder__is_active=True
    )
    if query:
        arsip_qs = arsip_qs.filter(
            Q(name__icontains=query) |
            Q(description__icontains=query) |
            Q(creator__icontains=query)
        )
    if tahun:
        arsip_qs = arsip_qs.filter(folder__year=tahun)
    if kategori:
        arsip_qs = arsip_qs.filter(file_type=kategori)

    wb = Workbook()
    ws = wb.active
    ws.title = "Daftar Arsip"

    headers = [
        'No', 'Nama File', 'Kode Referensi', 'Tahun', 'Tipe',
        'Ukuran (KB)', 'Pencipta', 'Deskripsi', 'Tanggal Upload', 'Folder'
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color='0058A8', end_color='0058A8', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for i, a in enumerate(arsip_qs, start=1):
        ws.append([
            i,
            a.name,
            a.reference_code or '-',
            a.folder.year if a.folder else '-',
            a.get_file_type_display(),
            round(a.size_bytes / 1024, 2),
            a.creator or '-',
            a.description or '-',
            a.uploaded_at.strftime('%d/%m/%Y %H:%M') if a.uploaded_at else '-',
            a.folder.name if a.folder else '-',
        ])

    widths = [5, 40, 20, 8, 10, 12, 20, 30, 18, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'daftar-arsip-{datetime.now().strftime("%Y%m%d-%H%M")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

@user_passes_test(is_admin)
def bulk_delete(request):
    """Bulk delete: hapus banyak arsip sekaligus"""
    if request.method == 'POST':
        arsip_ids = request.POST.getlist('arsip_ids')

        if not arsip_ids:
            messages.warning(request, 'Gak ada arsip yang dipilih.')
            return redirect('arsip:list')

        deleted = Arsip.objects.filter(pk__in=arsip_ids).update(is_active=False)

        log_activity(
            request.user, 'delete',
            f'Bulk delete {deleted} arsip',
            target_type='arsip', request=request
        )

        messages.success(request, f'{deleted} arsip berhasil dihapus.')
        return redirect('arsip:list')

    return redirect('arsip:list')


@user_passes_test(is_admin)
def bulk_move(request):
    """Bulk move: pindahkan banyak arsip ke folder lain"""
    if request.method == 'POST':
        arsip_ids = request.POST.getlist('arsip_ids')
        folder_id = request.POST.get('folder_id')

        if not arsip_ids:
            messages.warning(request, 'Gak ada arsip yang dipilih.')
            return redirect('arsip:list')

        if not folder_id:
            messages.error(request, 'Pilih folder tujuan dulu.')
            return redirect('arsip:list')

        try:
            target_folder = Folder.objects.get(pk=folder_id, is_active=True)
        except Folder.DoesNotExist:
            messages.error(request, 'Folder tujuan gak ditemukan.')
            return redirect('arsip:list')

        moved = Arsip.objects.filter(pk__in=arsip_ids).update(folder=target_folder)

        log_activity(
            request.user, 'move',
            f'Pindah {moved} arsip ke "{target_folder.name}"',
            target_type='arsip', target_id=target_folder.pk,
            request=request
        )

        messages.success(request, f'{moved} arsip berhasil dipindahkan ke "{target_folder.name}".')
        return redirect('arsip:folder', pk=target_folder.pk)

    return redirect('arsip:list')

@user_passes_test(is_admin)
def activity_log(request):
    """Halaman log aktivitas"""
    from django.core.paginator import Paginator
    
    # Filter
    action_filter = request.GET.get('action', '')
    user_filter = request.GET.get('user', '')
    page = request.GET.get('page', 1)
    
    logs = ActivityLog.objects.select_related('user').all()
    
    if action_filter:
        logs = logs.filter(action=action_filter)
    if user_filter:
        logs = logs.filter(user__username=user_filter)
    
    # Pagination
    paginator = Paginator(logs, 30)
    try:
        log_page = paginator.page(page)
    except:
        log_page = paginator.page(1)
    
    # User list buat filter
    from django.contrib.auth.models import User
    users = User.objects.filter(activities__isnull=False).distinct()
    
    context = {
        'log_list': log_page,
        'page_obj': log_page,
        'paginator': paginator,
        'action_filter': action_filter,
        'user_filter': user_filter,
        'action_choices': ActivityLog.ACTION_CHOICES,
        'users': users,
    }
    return render(request, 'arsip/activity_log.html', context)